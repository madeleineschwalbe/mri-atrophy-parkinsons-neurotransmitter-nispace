"""Generate appendix_results.xlsx — cortical volume (B) + desikanaseg combined (C)."""

import ast
import re
import numpy as np
import pandas as pd
from pathlib import Path
from scipy.stats import ttest_1samp
from statsmodels.stats.multitest import multipletests
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

RESULTS_DIR = Path("../../results")
DATA_DIR    = Path("../../data")
OUTPUT_FILE = RESULTS_DIR / "appendix_results.xlsx"

# ── Group sizes ──────────────────────────────────────────────────────────────
GROUP_SIZES = {
    "De Novo PD vs HC":           (558, 192),
    "Prodromal PD vs HC":         (279, 192),
    "De Novo PD vs Prodromal PD": (558, 279),
}

CONTRAST_ORDER_MAIN = [
    "De Novo PD vs HC",
    "Prodromal PD vs HC",
    "De Novo PD vs Prodromal PD",
]

# ── DK label → display name ──────────────────────────────────────────────────
DK_LABEL_MAP = {
    "bankssts": "Banks STS",
    "caudalanteriorcingulate": "Caudal anterior cingulate",
    "caudalmiddlefrontal": "Caudal middle frontal",
    "cuneus": "Cuneus",
    "entorhinal": "Entorhinal",
    "fusiform": "Fusiform",
    "inferiorparietal": "Inferior parietal",
    "inferiortemporal": "Inferior temporal",
    "isthmuscingulate": "Isthmus cingulate",
    "lateraloccipital": "Lateral occipital",
    "lateralorbitofrontal": "Lateral orbitofrontal",
    "lingual": "Lingual",
    "medialorbitofrontal": "Medial orbitofrontal",
    "middletemporal": "Middle temporal",
    "parahippocampal": "Parahippocampal",
    "paracentral": "Paracentral",
    "parsopercularis": "Pars opercularis",
    "parsorbitalis": "Pars orbitalis",
    "parstriangularis": "Pars triangularis",
    "pericalcarine": "Pericalcarine",
    "postcentral": "Postcentral",
    "posteriorcingulate": "Posterior cingulate",
    "precentral": "Precentral",
    "precuneus": "Precuneus",
    "rostralanteriorcingulate": "Rostral anterior cingulate",
    "rostralmiddlefrontal": "Rostral middle frontal",
    "superiorfrontal": "Superior frontal",
    "superiorparietal": "Superior parietal",
    "superiortemporal": "Superior temporal",
    "supramarginal": "Supramarginal",
    "frontalpole": "Frontal pole",
    "temporalpole": "Temporal pole",
    "transversetemporal": "Transverse temporal",
    "insula": "Insula",
}

ASEG_LABEL_MAP = {
    "Amygdala": "Amygdala",
    "Caudate": "Caudate",
    "Hippocampus": "Hippocampus",
    "Pallidum": "Pallidum",
    "Putamen": "Putamen",
    "Thalamus": "Thalamus",
    "Accumbens-area": "Accumbens area",
    "VentralDC": "Ventral DC",
    "Cerebellum-Cortex": "Cerebellum cortex",
}

# ── Reference map ordering ────────────────────────────────────────────────────
MAP_ORDER = [
    "Serotonin | 5HT1a", "Serotonin | 5HT1b", "Serotonin | 5HT2a",
    "Serotonin | 5HT4", "Serotonin | 5HTT",
    "Dopamine | D1", "Dopamine | D23", "Dopamine | DAT", "Dopamine | FDOPA",
    "GABA | GABAa",
    "Glutamate | mGluR5", "Glutamate | NMDA",
    "Noradrenaline/Acetylcholine | NET", "Noradrenaline/Acetylcholine | VAChT",
]
MAP_ORDER_IDX = {k: i for i, k in enumerate(MAP_ORDER)}

MAP_SHORT = {
    "Serotonin | 5HT1a": "5-HT1a",
    "Serotonin | 5HT1b": "5-HT1b",
    "Serotonin | 5HT2a": "5-HT2a",
    "Serotonin | 5HT4": "5-HT4",
    "Serotonin | 5HTT": "5-HTT",
    "Dopamine | D1": "D1",
    "Dopamine | D23": "D2/3",
    "Dopamine | DAT": "DAT",
    "Dopamine | FDOPA": "FDOPA",
    "GABA | GABAa": "GABA-A",
    "Glutamate | mGluR5": "mGluR5",
    "Glutamate | NMDA": "NMDA",
    "Noradrenaline/Acetylcholine | NET": "NET",
    "Noradrenaline/Acetylcholine | VAChT": "VAChT",
}

SYSTEM_DISPLAY = {
    "Serotonin": "Serotonin",
    "Dopamine": "Dopamine",
    "GABA": "GABA",
    "Glutamate": "Glutamate",
    "Noradrenaline/Acetylcholine": "NA/ACh",
}

CLINICAL_DISPLAY = {
    "moca_z": "MoCA",
    "updrs3_score_z": "UPDRS-III",
    "gds_z": "GDS",
}
CLINICAL_ORDER = ["moca_z", "updrs3_score_z", "gds_z"]

# ── Colors ───────────────────────────────────────────────────────────────────
HEADER_FILL  = PatternFill("solid", fgColor="2E4057")
NOM_FILL     = PatternFill("solid", fgColor="FFF3CD")
FDR_FILL     = PatternFill("solid", fgColor="D4EDDA")
ALT_FILL     = PatternFill("solid", fgColor="F8F9FA")
HEADER_FONT  = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
BODY_FONT    = Font(name="Calibri", size=10)
BOLD_FONT    = Font(name="Calibri", bold=True, size=10)
THIN_BORDER  = Border(
    bottom=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
)


# ── Helpers ──────────────────────────────────────────────────────────────────

def parse_parcel(parcel_str):
    """'hemi-L_lab-bankssts' → (region_display, hemi)"""
    m = re.match(r"hemi-([LR])_lab-(.+)", parcel_str)
    if not m:
        return parcel_str, ""
    hemi = m.group(1)
    label = m.group(2)
    display = DK_LABEL_MAP.get(label) or ASEG_LABEL_MAP.get(label) or label.capitalize()
    return display, hemi


def hedges_g(T, n1, n2, df):
    J = 1 - 3 / (4 * df - 1)
    return -T * np.sqrt(1 / n1 + 1 / n2) * J


def fdr_correct(p_values):
    reject, q, _, _ = multipletests(p_values, method="fdr_bh")
    return q


def parse_ref_map(s):
    """Parse NiSpace reference map string → (system_display, map_short, sort_key)"""
    try:
        tup = ast.literal_eval(s)
        system_raw = tup[0]
        target_raw = tup[1]
    except Exception:
        return "Unknown", s, 99
    m = re.search(r"target-([^_]+)", target_raw)
    target = m.group(1) if m else target_raw
    key = f"{system_raw} | {target}"
    system_display = SYSTEM_DISPLAY.get(system_raw, system_raw)
    map_short = MAP_SHORT.get(key, target)
    sort_key = MAP_ORDER_IDX.get(key, 99)
    return system_display, map_short, sort_key


def map_str_to_key(map_str):
    """'Glutamate | target-mGluR5_tracer-...' → 'Glutamate | mGluR5'"""
    parts = map_str.split(" | ")
    if len(parts) != 2:
        return None
    system = parts[0]
    m = re.search(r"target-([^_]+)", parts[1])
    if not m:
        return None
    return f"{system} | {m.group(1)}"


# ── Data loaders ─────────────────────────────────────────────────────────────

def load_parcelwise(csv_map: dict, contrast_order: list) -> pd.DataFrame:
    rows = []
    for contrast in contrast_order:
        path = csv_map[contrast]
        df = pd.read_csv(path)
        n1, n2 = GROUP_SIZES[contrast]
        df["contrast"] = contrast
        df["n1"] = n1
        df["n2"] = n2
        df["hedges_g"] = df.apply(
            lambda r: hedges_g(r["Tvalue"], n1, n2, r["df"]), axis=1
        )
        df["q_fdr"] = fdr_correct(df["pvalue"].values)
        rows.append(df)

    combined = pd.concat(rows, ignore_index=True)
    parsed = combined["parcel"].apply(parse_parcel)
    combined["Region"] = [p[0] for p in parsed]
    combined["Hemisphere"] = [p[1] for p in parsed]

    out = pd.DataFrame({
        "Contrast":      combined["contrast"],
        "Region":        combined["Region"],
        "Hemisphere":    combined["Hemisphere"],
        "n₁":            combined["n1"],
        "n₂":            combined["n2"],
        "T-statistic":   combined["Tvalue"].round(3),
        "df":            combined["df"].round(1),
        "p (uncorr.)":   combined["pvalue"].round(4),
        "Hedges' g":     combined["hedges_g"].round(3),
        "q (FDR)":       combined["q_fdr"].round(4),
        "Sig. (p<0.05)": combined["pvalue"].lt(0.05).map({True: "★", False: ""}),
        "Sig. (FDR)":    combined["q_fdr"].lt(0.05).map({True: "★★", False: ""}),
    })

    contrast_cat = pd.Categorical(out["Contrast"], categories=contrast_order, ordered=True)
    out["_contrast_order"] = contrast_cat
    out = out.sort_values(["_contrast_order", "Hedges' g"],
                          key=lambda s: s if s.name != "Hedges' g" else s.abs(),
                          ascending=[True, False]).drop(columns="_contrast_order")
    return out.reset_index(drop=True)


def load_parcelwise_desikanaseg(contrast_order: list) -> pd.DataFrame:
    """Combine cortical + subcortical parcelwise results for Desikanaseg atlas."""
    cortical_files = {
        "De Novo PD vs HC":           RESULTS_DIR / "de novo pd vs hc_parcelwise_ttest_volume_cortical.csv",
        "Prodromal PD vs HC":         RESULTS_DIR / "prodromal pd vs hc_parcelwise_ttest_volume_cortical.csv",
        "De Novo PD vs Prodromal PD": RESULTS_DIR / "de novo pd vs prodromal pd_parcelwise_ttest_volume_cortical.csv",
    }
    subcortical_files = {
        "De Novo PD vs HC":           RESULTS_DIR / "de_novo_pd_vs_hc_parcelwise_ttest_volume_subcortical.csv",
        "Prodromal PD vs HC":         RESULTS_DIR / "prodromal_pd_vs_hc_parcelwise_ttest_volume_subcortical.csv",
        "De Novo PD vs Prodromal PD": RESULTS_DIR / "de_novo_pd_vs_prodromal_pd_parcelwise_ttest_volume_subcortical.csv",
    }

    rows = []
    for contrast in contrast_order:
        n1, n2 = GROUP_SIZES[contrast]
        df_cort = pd.read_csv(cortical_files[contrast])
        df_cort["struct_type"] = "Cortical"
        df_sub = pd.read_csv(subcortical_files[contrast])
        df_sub["struct_type"] = "Subcortical"
        combined = pd.concat([df_cort, df_sub], ignore_index=True)
        combined["contrast"] = contrast
        combined["n1"] = n1
        combined["n2"] = n2
        combined["hedges_g"] = combined.apply(
            lambda r: hedges_g(r["Tvalue"], n1, n2, r["df"]), axis=1
        )
        # FDR within contrast across all parcels (cortical + subcortical), skip NaN p-values
        valid = combined["pvalue"].notna()
        combined["q_fdr"] = np.nan
        if valid.sum() > 0:
            combined.loc[valid, "q_fdr"] = fdr_correct(combined.loc[valid, "pvalue"].values)
        rows.append(combined)

    all_data = pd.concat(rows, ignore_index=True)
    parsed = all_data["parcel"].apply(parse_parcel)
    all_data["Region"] = [p[0] for p in parsed]
    all_data["Hemisphere"] = [p[1] for p in parsed]

    out = pd.DataFrame({
        "Contrast":      all_data["contrast"],
        "Type":          all_data["struct_type"],
        "Region":        all_data["Region"],
        "Hemisphere":    all_data["Hemisphere"],
        "n₁":            all_data["n1"],
        "n₂":            all_data["n2"],
        "T-statistic":   all_data["Tvalue"].round(3),
        "df":            all_data["df"].round(1),
        "p (uncorr.)":   all_data["pvalue"].round(4),
        "Hedges' g":     all_data["hedges_g"].round(3),
        "q (FDR)":       all_data["q_fdr"].round(4),
        "Sig. (p<0.05)": all_data["pvalue"].lt(0.05).map({True: "★", False: ""}),
        "Sig. (FDR)":    all_data["q_fdr"].lt(0.05).map({True: "★★", False: ""}),
    })

    type_cat = pd.Categorical(out["Type"], categories=["Cortical", "Subcortical"], ordered=True)
    contrast_cat = pd.Categorical(out["Contrast"], categories=contrast_order, ordered=True)
    out["_contrast_order"] = contrast_cat
    out["_type_order"] = type_cat
    out = out.sort_values(
        ["_contrast_order", "_type_order", "Hedges' g"],
        key=lambda s: s if s.name != "Hedges' g" else s.abs(),
        ascending=[True, True, False],
    ).drop(columns=["_contrast_order", "_type_order"])
    return out.reset_index(drop=True)


def load_colocalization(csv_path: str, contrast_order: list) -> pd.DataFrame:
    df = pd.read_csv(csv_path)
    parsed = df["reference_map"].apply(parse_ref_map)
    df["System"]   = [p[0] for p in parsed]
    df["Map"]      = [p[1] for p in parsed]
    df["_map_ord"] = [p[2] for p in parsed]

    contrast_cat = pd.Categorical(df["contrast"], categories=contrast_order, ordered=True)
    df["_contrast_ord"] = contrast_cat
    df = df.sort_values(["_contrast_ord", "_map_ord"])

    out = pd.DataFrame({
        "Contrast":      df["contrast"].values,
        "System":        df["System"].values,
        "Map":           df["Map"].values,
        "ρ (Spearman)":  df["rho"].round(3).values,
        "p (perm.)":     df["p"].round(4).values,
        "q (FDR)":       df["q"].round(4).values,
        "Sig. (p<0.05)": (df["p"] < 0.05).map({True: "★", False: ""}).values,
        "Sig. (FDR)":    (df["q"] < 0.05).map({True: "★★", False: ""}).values,
    })
    return out.reset_index(drop=True)


def load_single_subject_analysis(csv_path: str, contrast_order: list) -> pd.DataFrame:
    """
    Per-subject Fisher's z colocalization scores → one-sample t-test vs. 0 per map/contrast.
    Shows whether subjects' individual NiSpace scores systematically differ from zero.
    """
    df = pd.read_csv(csv_path)
    df["fz"] = np.arctanh(df["colocalization"].clip(-0.9999, 0.9999))
    df["map_key"] = df["map"].apply(map_str_to_key)
    df = df.dropna(subset=["map_key"])

    rows = []
    for contrast in contrast_order:
        df_c = df[df["contrast"] == contrast]
        for map_key in MAP_ORDER:
            fz_vals = df_c.loc[df_c["map_key"] == map_key, "fz"].dropna().values
            if len(fz_vals) < 2:
                continue
            system_raw, target = map_key.split(" | ")
            t_stat, p_val = ttest_1samp(fz_vals, popmean=0)
            rows.append({
                "Contrast":         contrast,
                "System":           SYSTEM_DISPLAY.get(system_raw, system_raw),
                "Map":              MAP_SHORT.get(map_key, target),
                "_map_ord":         MAP_ORDER_IDX.get(map_key, 99),
                "N":                len(fz_vals),
                "Mean Fisher's z":  round(float(np.mean(fz_vals)), 3),
                "SD":               round(float(np.std(fz_vals, ddof=1)), 3),
                "Median Fisher's z":round(float(np.median(fz_vals)), 3),
                "Q1":               round(float(np.percentile(fz_vals, 25)), 3),
                "Q3":               round(float(np.percentile(fz_vals, 75)), 3),
                "t-statistic":      round(float(t_stat), 3),
                "df":               int(len(fz_vals) - 1),
                "p (one-sample)":   round(float(p_val), 6),
            })

    out = pd.DataFrame(rows)

    # FDR correction within each contrast (across 14 maps)
    out["q (FDR)"] = np.nan
    for contrast in contrast_order:
        mask = out["Contrast"] == contrast
        if mask.sum() > 0:
            out.loc[mask, "q (FDR)"] = fdr_correct(out.loc[mask, "p (one-sample)"].values).round(6)

    out["Sig. (p<0.05)"] = out["p (one-sample)"].lt(0.05).map({True: "★", False: ""})
    out["Sig. (FDR)"]    = out["q (FDR)"].lt(0.05).map({True: "★★", False: ""})

    contrast_cat = pd.Categorical(out["Contrast"], categories=contrast_order, ordered=True)
    out["_contrast_ord"] = contrast_cat
    out = out.sort_values(["_contrast_ord", "_map_ord"]).drop(columns=["_map_ord", "_contrast_ord"])

    return out[["Contrast", "System", "Map", "N",
                "Mean Fisher's z", "SD", "Median Fisher's z", "Q1", "Q3",
                "t-statistic", "df", "p (one-sample)", "q (FDR)",
                "Sig. (p<0.05)", "Sig. (FDR)"]].reset_index(drop=True)


def load_clinical_corr(csv_path: str, contrast_order: list) -> pd.DataFrame:
    df = pd.read_csv(csv_path)

    fdr_col = next((c for c in df.columns if "spearman" in c.lower() and "fdr" in c.lower()), None)
    if fdr_col is None:
        raise ValueError(f"No Spearman FDR column found in {csv_path}")

    def split_map(s):
        parts = s.split(" | ")
        if len(parts) == 2:
            system_raw, target = parts
            key = s
            system_disp = SYSTEM_DISPLAY.get(system_raw, system_raw)
            map_short = MAP_SHORT.get(key, target)
            sort_k = MAP_ORDER_IDX.get(key, 99)
        else:
            system_disp, map_short, sort_k = "Unknown", s, 99
        return system_disp, map_short, sort_k

    parsed = df["map"].apply(split_map)
    df["System"]   = [p[0] for p in parsed]
    df["Map"]      = [p[1] for p in parsed]
    df["_map_ord"] = [p[2] for p in parsed]

    df["Clinical Variable"] = df["clinical_var"].map(CLINICAL_DISPLAY).fillna(df["clinical_var"])
    clin_cat = pd.Categorical(df["clinical_var"], categories=CLINICAL_ORDER, ordered=True)
    df["_clin_ord"] = clin_cat

    contrast_cat = pd.Categorical(df["contrast"], categories=contrast_order, ordered=True)
    df["_contrast_ord"] = contrast_cat
    df = df.sort_values(["_contrast_ord", "_map_ord", "_clin_ord"])

    out = pd.DataFrame({
        "Contrast":          df["contrast"].values,
        "System":            df["System"].values,
        "Map":               df["Map"].values,
        "Clinical Variable": df["Clinical Variable"].values,
        "Spearman r":        df["spearman_r"].round(3).values,
        "p (uncorr.)":       df["spearman_p"].round(4).values,
        "q (FDR)":           df[fdr_col].round(4).values,
        "N":                 df["n"].astype(int).values,
        "Sig. (p<0.05)":     (df["spearman_p"] < 0.05).map({True: "★", False: ""}).values,
        "Sig. (FDR)":        (df[fdr_col] < 0.05).map({True: "★★", False: ""}).values,
    })
    return out.reset_index(drop=True)


# ── Excel writer ─────────────────────────────────────────────────────────────

def write_sheet(ws, df: pd.DataFrame, title: str):
    ws.append([title])
    title_cell = ws.cell(row=1, column=1)
    title_cell.font = Font(name="Calibri", bold=True, size=13, color="2E4057")
    ws.row_dimensions[1].height = 20

    ws.append([])

    header_row = 3
    ws.append(list(df.columns))
    for col_idx in range(1, len(df.columns) + 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 30

    for row_idx, (_, row) in enumerate(df.iterrows(), start=4):
        is_alt  = (row_idx - 4) % 2 == 1
        nom_sig = row.get("Sig. (p<0.05)", "") == "★"
        fdr_sig = row.get("Sig. (FDR)", "") == "★★"

        if fdr_sig:
            fill = FDR_FILL
        elif nom_sig:
            fill = NOM_FILL
        elif is_alt:
            fill = ALT_FILL
        else:
            fill = None

        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = BOLD_FONT if fdr_sig else BODY_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if fill:
                cell.fill = fill

    for col_idx, col_name in enumerate(df.columns, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = max(
            len(str(col_name)),
            df.iloc[:, col_idx - 1].astype(str).str.len().max()
        )
        ws.column_dimensions[col_letter].width = min(max_len + 3, 35)

    ws.freeze_panes = ws.cell(row=4, column=1)


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    print("Building appendix XLSX...")

    # ── B1: Cortical Volume parcelwise (DK 68) ────────────────────────────────
    print("  B1: Cortical volume parcelwise")
    b1 = load_parcelwise(
        csv_map={
            "De Novo PD vs HC":           RESULTS_DIR / "de novo pd vs hc_parcelwise_ttest_volume_cortical.csv",
            "Prodromal PD vs HC":         RESULTS_DIR / "prodromal pd vs hc_parcelwise_ttest_volume_cortical.csv",
            "De Novo PD vs Prodromal PD": RESULTS_DIR / "de novo pd vs prodromal pd_parcelwise_ttest_volume_cortical.csv",
        },
        contrast_order=CONTRAST_ORDER_MAIN,
    )

    # ── B2: Cortical Volume colocalization ────────────────────────────────────
    print("  B2: Cortical volume colocalization")
    b2 = load_colocalization(
        RESULTS_DIR / "nispace_group_comparison_results_volumes_cortical.csv",
        contrast_order=CONTRAST_ORDER_MAIN,
    )

    # ── B3: Cortical Volume clinical correlation ──────────────────────────────
    print("  B3: Cortical volume clinical corr")
    b3 = load_clinical_corr(
        RESULTS_DIR / "clinical_correlation_volume_cortical.csv",
        contrast_order=CONTRAST_ORDER_MAIN,
    )

    # ── C1: Desikanaseg parcelwise (cortical + subcortical combined) ──────────
    print("  C1: Desikanaseg parcelwise (combined)")
    c1 = load_parcelwise_desikanaseg(contrast_order=CONTRAST_ORDER_MAIN)

    # ── C2: Desikanaseg colocalization ────────────────────────────────────────
    print("  C2: Desikanaseg colocalization")
    c2 = load_colocalization(
        RESULTS_DIR / "nispace_group_comparison_results_volumes_desikanaseg.csv",
        contrast_order=CONTRAST_ORDER_MAIN,
    )

    # ── C3: Desikanaseg single-subject group analysis ─────────────────────────
    print("  C3: Desikanaseg single-subject group analysis")
    c3 = load_single_subject_analysis(
        DATA_DIR / "merged_df_volume_desikanaseg.csv",
        contrast_order=CONTRAST_ORDER_MAIN,
    )

    # ── C4: Desikanaseg clinical correlation ──────────────────────────────────
    print("  C4: Desikanaseg clinical corr")
    c4 = load_clinical_corr(
        RESULTS_DIR / "clinical_correlation_volume_desikanaseg.csv",
        contrast_order=CONTRAST_ORDER_MAIN,
    )

    # ── Write Excel ───────────────────────────────────────────────────────────
    sheets = [
        ("B1 Cort Vol Parcelwise",
         b1, "Appendix B1 — Cortical Volume Parcelwise Results (Desikan-Killiany, 68 parcels)"),
        ("B2 Cort Vol Colocaliz",
         b2, "Appendix B2 — Cortical Volume NiSpace Colocalization"),
        ("B3 Cort Vol Clin Corr",
         b3, "Appendix B3 — Cortical Volume Clinical Correlations"),
        ("C1 Desikanaseg Parcelwise",
         c1, "Appendix C1 — Desikanaseg Volume Parcelwise Results (Cortical + Subcortical)"),
        ("C2 Desikanaseg Colocaliz",
         c2, "Appendix C2 — Desikanaseg Volume NiSpace Colocalization"),
        ("C3 Desikanaseg Single-Subject",
         c3, "Appendix C3 — Desikanaseg Volume Single-Subject Group Analysis (one-sample t-test of Fisher's z vs. 0)"),
        ("C4 Desikanaseg Clin Corr",
         c4, "Appendix C4 — Desikanaseg Volume Clinical Correlations"),
    ]

    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        for sheet_name, df, title in sheets:
            df.to_excel(writer, sheet_name=sheet_name, index=False)

    wb = load_workbook(OUTPUT_FILE)
    for sheet_name, df, title in sheets:
        ws = wb[sheet_name]
        ws.delete_rows(1, ws.max_row)
        write_sheet(ws, df, title)

    # ── README ────────────────────────────────────────────────────────────────
    ws_readme = wb.create_sheet("README", 0)
    readme_lines = [
        ("APPENDIX RESULTS", Font(name="Calibri", bold=True, size=14, color="2E4057")),
        ("", None),
        ("Sheet overview:", Font(name="Calibri", bold=True, size=11)),
        ("B1 Cort Vol Parcelwise         — Parcelwise t-test results (cortical volume, DK 68 parcels), 3 main contrasts", None),
        ("B2 Cort Vol Colocaliz          — NiSpace colocalization for cortical volume × 3 main contrasts", None),
        ("B3 Cort Vol Clin Corr          — Clinical correlations for cortical volume × 3 main contrasts", None),
        ("C1 Desikanaseg Parcelwise      — Parcelwise t-test results (cortical + subcortical combined), 3 main contrasts", None),
        ("C2 Desikanaseg Colocaliz       — NiSpace colocalization for desikanaseg volume × 3 main contrasts", None),
        ("C3 Desikanaseg Single-Subject  — Per-subject Fisher's z colocalization (one-sample t-test vs. 0) × 3 main contrasts", None),
        ("C4 Desikanaseg Clin Corr       — Clinical correlations for desikanaseg volume × 3 main contrasts", None),
        ("", None),
        ("Significance coding:", Font(name="Calibri", bold=True, size=11)),
        ("★   Nominally significant (p < 0.05, uncorrected)", None),
        ("★★  FDR-significant (Benjamini-Hochberg q < 0.05)", None),
        ("", None),
        ("Row highlighting:", Font(name="Calibri", bold=True, size=11)),
        ("Light yellow = p < 0.05 (nominal)", None),
        ("Light green  = q < 0.05 (FDR)", None),
        ("", None),
        ("NiSpace sign convention:", Font(name="Calibri", bold=True, size=11)),
        ("Hedges' g = −T × √(1/n₁ + 1/n₂) × J", None),
        ("Positive g = greater value in first-named group (e.g. De Novo PD > HC)", None),
        ("FDR correction applied within each contrast.", None),
        ("", None),
        ("C3 note:", Font(name="Calibri", bold=True, size=11)),
        ("Fisher's z = arctanh(Spearman ρ) of per-subject deviation map vs. reference map.", None),
        ("One-sample t-test tests whether mean Fisher's z differs from zero (i.e., systematic alignment with the reference map).", None),
    ]

    for i, (text, font) in enumerate(readme_lines, start=1):
        cell = ws_readme.cell(row=i, column=1, value=text)
        cell.font = font if font else BODY_FONT
    ws_readme.column_dimensions["A"].width = 100

    wb.save(OUTPUT_FILE)
    print(f"\nSaved → {OUTPUT_FILE.resolve()}")

    for sheet_name, df, _ in sheets:
        print(f"  {sheet_name:<35} {len(df):>4} rows")


if __name__ == "__main__":
    main()
